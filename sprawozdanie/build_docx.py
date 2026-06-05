#!/usr/bin/env python3
"""Generuje sformatowany plik .docx ze sprawozdania w formacie Markdown.

Obsługuje: nagłówki (#, ##, ###), akapity z pogrubieniem (**...**), listy
punktowane i numerowane, tabele Markdown (| ... |) oraz bloki kodu (```).

Formatowanie pracy:
- czcionka tekstu: Times New Roman 12, interlinia 1.5, justowanie
- marginesy 2.5 cm, numeracja stron w stopce
- kod: Consolas 10

Użycie:
    pip install python-docx
    python3 build_docx.py            # sprawozdanie.md -> sprawozdanie.docx
"""
import os
import re
import sys

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
from docx.shared import Pt, Cm, RGBColor

BASE = os.path.dirname(os.path.abspath(__file__))
SRC = os.path.join(BASE, "sprawozdanie.md")
OUT = os.path.join(BASE, "sprawozdanie.docx")


def set_cell_bg(cell, color):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:fill"), color)
    tcPr.append(shd)


def add_page_numbers(section):
    footer = section.footer
    p = footer.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run()
    fldChar1 = OxmlElement("w:fldChar"); fldChar1.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText"); instr.set(qn("xml:space"), "preserve"); instr.text = "PAGE"
    fldChar2 = OxmlElement("w:fldChar"); fldChar2.set(qn("w:fldCharType"), "end")
    run._r.append(fldChar1); run._r.append(instr); run._r.append(fldChar2)


def add_runs_with_bold(paragraph, text):
    """Dodaje tekst, interpretując **pogrubienie** oraz `kod`."""
    parts = re.split(r"(\*\*.+?\*\*|`.+?`)", text)
    for part in parts:
        if not part:
            continue
        if part.startswith("**") and part.endswith("**"):
            r = paragraph.add_run(part[2:-2]); r.bold = True
        elif part.startswith("`") and part.endswith("`"):
            r = paragraph.add_run(part[1:-1]); r.font.name = "Consolas"; r.font.size = Pt(10)
        else:
            paragraph.add_run(part)


def style_normal(doc):
    st = doc.styles["Normal"]
    st.font.name = "Times New Roman"
    st.font.size = Pt(12)
    pf = st.paragraph_format
    pf.line_spacing = 1.5
    pf.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    pf.space_after = Pt(6)


def main():
    src = sys.argv[1] if len(sys.argv) > 1 else SRC
    out = sys.argv[2] if len(sys.argv) > 2 else OUT
    with open(src, encoding="utf-8") as f:
        lines = f.read().splitlines()

    doc = Document()
    for s in doc.sections:
        s.top_margin = s.bottom_margin = s.left_margin = s.right_margin = Cm(2.5)
        add_page_numbers(s)
    style_normal(doc)

    i = 0
    n = len(lines)
    while i < n:
        line = lines[i]

        # blok kodu
        if line.strip().startswith("```"):
            i += 1
            code = []
            while i < n and not lines[i].strip().startswith("```"):
                code.append(lines[i]); i += 1
            i += 1
            p = doc.add_paragraph()
            p.paragraph_format.line_spacing = 1.0
            p.paragraph_format.left_indent = Cm(0.5)
            r = p.add_run("\n".join(code))
            r.font.name = "Consolas"; r.font.size = Pt(10)
            r.font.color.rgb = RGBColor(0x22, 0x22, 0x22)
            continue

        # tabela Markdown
        if line.strip().startswith("|") and i + 1 < n and re.match(r"^\s*\|[\s:|-]+\|\s*$", lines[i + 1]):
            header = [c.strip() for c in line.strip().strip("|").split("|")]
            i += 2
            rows = []
            while i < n and lines[i].strip().startswith("|"):
                rows.append([c.strip() for c in lines[i].strip().strip("|").split("|")])
                i += 1
            tbl = doc.add_table(rows=1, cols=len(header))
            tbl.style = "Table Grid"
            tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
            for j, h in enumerate(header):
                cell = tbl.rows[0].cells[j]
                cell.paragraphs[0].text = ""
                run = cell.paragraphs[0].add_run(h.replace("**", "")); run.bold = True
                run.font.size = Pt(10); run.font.name = "Times New Roman"
                set_cell_bg(cell, "D9E2F3")
            for row in rows:
                cells = tbl.add_row().cells
                for j in range(len(header)):
                    val = row[j] if j < len(row) else ""
                    cells[j].paragraphs[0].text = ""
                    add_runs_with_bold(cells[j].paragraphs[0], val)
                    for rn in cells[j].paragraphs[0].runs:
                        rn.font.size = Pt(10); rn.font.name = "Times New Roman"
            doc.add_paragraph()
            continue

        # nagłówki
        if line.startswith("### "):
            h = doc.add_heading(level=3); add_runs_with_bold(h, line[4:])
        elif line.startswith("## "):
            h = doc.add_heading(level=2); add_runs_with_bold(h, line[3:])
        elif line.startswith("# "):
            h = doc.add_heading(level=1); add_runs_with_bold(h, line[2:])
        elif line.strip() == "---":
            pass
        elif re.match(r"^\s*[-*] ", line):
            p = doc.add_paragraph(style="List Bullet")
            add_runs_with_bold(p, re.sub(r"^\s*[-*] ", "", line))
        elif re.match(r"^\s*\d+\. ", line):
            p = doc.add_paragraph(style="List Number")
            add_runs_with_bold(p, re.sub(r"^\s*\d+\. ", "", line))
        elif line.startswith("> "):
            p = doc.add_paragraph(); p.paragraph_format.left_indent = Cm(0.8)
            add_runs_with_bold(p, line[2:]);
            for rn in p.runs: rn.italic = True
        elif line.strip() == "":
            pass
        else:
            p = doc.add_paragraph(); add_runs_with_bold(p, line)
        i += 1

    doc.save(out)
    print(f"Zapisano: {out}")


if __name__ == "__main__":
    main()
