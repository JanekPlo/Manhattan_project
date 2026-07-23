#!/usr/bin/env python3
"""Generuje sformatowany plik .docx z projekt-smartdorm.md.

Formatowanie zgodne z warunkami zaliczenia:
- czcionka: Times New Roman 12 pkt
- interlinia: 1,5
- marginesy normalne (2,5 cm) wg standardowego szablonu Word
- strona tytułowa z nr albumu oraz imieniem i nazwiskiem studenta
- numeracja stron (w stopce)
- obsługa obrazów markdown ![podpis](ścieżka) — załączniki (diagramy)

Użycie:
    pip install python-docx
    python generate_diagrams.py     # najpierw diagramy
    python build_smartdorm.py       # potem dokument
"""
import os
import re
import sys

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
from docx.shared import Pt, Cm, RGBColor

FONT_NAME = "Times New Roman"
FONT_SIZE = 12

# --- Dane strony tytułowej (edytuj wedle potrzeb) ---
TP_TITLE = "Projekt systemu informatycznego\nSmartDorm"
TP_SUBJECT = ("System automatyzacji procesów kwaterowania, opłat "
              "oraz zgłaszania usterek w akademiku")
TP_KIND = "Dokument zaliczeniowy — projekt systemu informatycznego"
TP_ALBUM = "Nr albumu: ................"   # <-- uzupełnij numer albumu
TP_AUTHOR = "Jan Płoński"
TP_DATE = "Lipiec 2026"


def _force_font(style):
    """Ustawia Times New Roman dla wszystkich wariantów (ascii/hAnsi/cs/eastAsia)."""
    style.font.name = FONT_NAME
    rpr = style.element.get_or_add_rPr()
    rfonts = rpr.get_or_add_rFonts()
    for attr in ("w:ascii", "w:hAnsi", "w:cs", "w:eastAsia"):
        rfonts.set(qn(attr), FONT_NAME)


def set_base_style(document):
    """Czcionka, rozmiar, interlinia 1.5 i justowanie dla stylu Normal."""
    style = document.styles["Normal"]
    style.font.size = Pt(FONT_SIZE)
    _force_font(style)
    pf = style.paragraph_format
    pf.line_spacing = 1.5
    pf.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    pf.space_after = Pt(6)
    for name in ("Title", "Heading 1", "Heading 2", "Heading 3",
                 "List Bullet", "List Number"):
        try:
            _force_font(document.styles[name])
        except KeyError:
            pass


def set_margins(document):
    for section in document.sections:
        section.top_margin = Cm(2.5)
        section.bottom_margin = Cm(2.5)
        section.left_margin = Cm(2.5)
        section.right_margin = Cm(2.5)


def add_page_numbers(document):
    """Numeruje strony w stopce; strona tytułowa pozostaje bez numeru."""
    for section in document.sections:
        section.different_first_page_header_footer = True
        sectPr = section._sectPr
        pg = OxmlElement("w:pgNumType")
        pg.set(qn("w:start"), "0")
        sectPr.append(pg)

        p = section.footer.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = p.add_run()
        fld_begin = OxmlElement("w:fldChar")
        fld_begin.set(qn("w:fldCharType"), "begin")
        instr = OxmlElement("w:instrText")
        instr.set(qn("xml:space"), "preserve")
        instr.text = "PAGE"
        fld_end = OxmlElement("w:fldChar")
        fld_end.set(qn("w:fldCharType"), "end")
        run._r.append(fld_begin)
        run._r.append(instr)
        run._r.append(fld_end)


def add_title_page(document):
    """Tworzy stronę tytułową zakończoną podziałem strony."""
    def centered(text, size, bold=False, italic=False, space_before=0, space_after=0):
        p = document.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_before = Pt(space_before)
        p.paragraph_format.space_after = Pt(space_after)
        p.paragraph_format.line_spacing = 1.0
        for j, line in enumerate(text.split("\n")):
            run = p.add_run(line)
            run.font.name = FONT_NAME
            run.font.size = Pt(size)
            run.bold = bold
            run.italic = italic
            if j < len(text.split("\n")) - 1:
                run.add_break()
        return p

    for _ in range(4):
        document.add_paragraph()
    centered(TP_TITLE, 20, bold=True, space_after=14)
    centered(TP_SUBJECT, 13, italic=True, space_after=30)
    centered(TP_KIND, 13, space_after=60)
    for _ in range(4):
        document.add_paragraph()
    centered(TP_AUTHOR, 14, bold=True, space_after=6)
    centered(TP_ALBUM, 12, space_after=6)
    centered(TP_DATE, 12)

    br = document.add_paragraph()
    br.add_run().add_break(WD_BREAK.PAGE)


def add_runs_with_bold(paragraph, text):
    """Dodaje tekst do akapitu, interpretując **pogrubienie** i `kod`."""
    for part in re.split(r"(\*\*.+?\*\*|`[^`]+`)", text):
        if not part:
            continue
        if part.startswith("**") and part.endswith("**"):
            for sub in re.split(r"(`[^`]+`)", part[2:-2]):
                if not sub:
                    continue
                if sub.startswith("`") and sub.endswith("`"):
                    run = paragraph.add_run(sub[1:-1])
                    run.font.name = "Courier New"
                    run.font.size = Pt(FONT_SIZE - 1)
                else:
                    run = paragraph.add_run(sub)
                run.bold = True
        elif part.startswith("`") and part.endswith("`"):
            run = paragraph.add_run(part[1:-1])
            run.font.name = "Courier New"
            run.font.size = Pt(FONT_SIZE - 1)
        else:
            paragraph.add_run(part)


def render_table(document, rows):
    """rows: lista list komórek (już bez znaków |)."""
    if not rows:
        return
    body = [r for r in rows if not re.match(r"^\s*:?-{2,}", r[0])]
    if not body:
        return
    cols = max(len(r) for r in body)
    table = document.add_table(rows=0, cols=cols)
    table.style = "Light Grid Accent 1"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    for r_idx, row in enumerate(body):
        cells = table.add_row().cells
        for c_idx in range(cols):
            txt = row[c_idx] if c_idx < len(row) else ""
            cell_p = cells[c_idx].paragraphs[0]
            cell_p.alignment = WD_ALIGN_PARAGRAPH.LEFT
            add_runs_with_bold(cell_p, txt.strip())
            if r_idx == 0:
                for run in cell_p.runs:
                    run.bold = True


def render_image(document, alt, path, base_dir):
    """Wstawia wyśrodkowany obraz na szerokość strony + podpis kursywą."""
    full = path if os.path.isabs(path) else os.path.join(base_dir, path)
    if not os.path.exists(full):
        print(f"UWAGA: brak pliku obrazu: {full}", file=sys.stderr)
        return
    p = document.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.line_spacing = 1.0
    # 16 cm = szerokość strony A4 minus marginesy 2 x 2,5 cm
    p.add_run().add_picture(full, width=Cm(16))
    if alt:
        cap = document.add_paragraph()
        cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
        cap.paragraph_format.line_spacing = 1.0
        run = cap.add_run(alt)
        run.italic = True
        run.font.size = Pt(10)


def parse_markdown(document, lines, base_dir="."):
    """Parser scalający miękko zawinięte linie w płynne akapity."""
    i = 0
    in_code = False
    code_buffer = []
    table_buffer = []
    quote_buffer = []
    block = None

    def flush_table():
        if table_buffer:
            render_table(document, list(table_buffer))
            table_buffer.clear()

    def flush_code():
        if code_buffer:
            p = document.add_paragraph()
            p.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.LEFT
            run = p.add_run("\n".join(code_buffer))
            run.font.name = "Courier New"
            run.font.size = Pt(9)
            code_buffer.clear()

    def flush_quote():
        if quote_buffer:
            p = document.add_paragraph()
            p.paragraph_format.left_indent = Cm(0.75)
            add_runs_with_bold(p, " ".join(quote_buffer))
            for run in p.runs:
                run.italic = True
                run.font.color.rgb = RGBColor(0x55, 0x55, 0x55)
            quote_buffer.clear()

    def flush_block():
        nonlocal block
        if not block:
            return
        kind, parts = block
        text = " ".join(parts).strip()
        block = None
        if not text:
            return
        if kind == "bullet":
            p = document.add_paragraph(style="List Bullet")
        elif kind.startswith("number:"):
            # literalna numeracja: styl List Number kontynuowałby licznik
            # między osobnymi listami w dokumencie
            p = document.add_paragraph()
            p.paragraph_format.left_indent = Cm(0.75)
            p.paragraph_format.first_line_indent = Cm(-0.5)
            text = f"{kind.split(':', 1)[1]}. {text}"
        else:
            p = document.add_paragraph()
        add_runs_with_bold(p, text)

    def flush_all():
        flush_block()
        flush_quote()
        flush_table()

    while i < len(lines):
        raw = lines[i].rstrip("\n")
        stripped = raw.strip()

        if stripped.startswith("```"):
            flush_all()
            if in_code:
                flush_code()
                in_code = False
            else:
                in_code = True
            i += 1
            continue
        if in_code:
            code_buffer.append(raw)
            i += 1
            continue

        # obrazy ![podpis](ścieżka)
        m_img = re.match(r"^!\[(.*?)\]\((.+?)\)\s*$", stripped)
        if m_img:
            flush_all()
            render_image(document, m_img.group(1), m_img.group(2), base_dir)
            i += 1
            continue

        if stripped.startswith("|") and stripped.endswith("|"):
            flush_block()
            flush_quote()
            cells = [c.strip() for c in stripped.strip("|").split("|")]
            table_buffer.append(cells)
            i += 1
            continue
        flush_table()

        if not stripped:
            flush_block()
            flush_quote()
            i += 1
            continue

        if re.match(r"^-{3,}$", stripped):
            flush_all()
            i += 1
            continue

        m_head = re.match(r"^(#{1,3})\s+(.*)", stripped)
        if m_head:
            flush_all()
            level = len(m_head.group(1))  # # -> Heading 1, ## -> Heading 2
            h = document.add_heading(level=level)
            # załączniki zaczynają się od nowej strony
            if level == 1 and m_head.group(2).startswith("Załącznik"):
                h.paragraph_format.page_break_before = True
            add_runs_with_bold(h, m_head.group(2))
            i += 1
            continue

        if stripped.startswith(">"):
            flush_block()
            quote_buffer.append(stripped.lstrip(">").strip())
            i += 1
            continue
        flush_quote()

        m_bullet = re.match(r"^[-*]\s+(.*)", stripped)
        m_num = re.match(r"^(\d+)\.\s+(.*)", stripped)
        if m_bullet:
            flush_block()
            block = ("bullet", [m_bullet.group(1)])
            i += 1
            continue
        if m_num:
            flush_block()
            block = (f"number:{m_num.group(1)}", [m_num.group(2)])
            i += 1
            continue

        if block:
            block[1].append(stripped)
        else:
            block = ("p", [stripped])
        i += 1

    flush_all()
    flush_code()


def main():
    src = sys.argv[1] if len(sys.argv) > 1 else "projekt-smartdorm.md"
    out = sys.argv[2] if len(sys.argv) > 2 else "projekt-smartdorm.docx"

    with open(src, encoding="utf-8") as f:
        lines = f.readlines()

    document = Document()
    set_base_style(document)
    set_margins(document)
    add_page_numbers(document)
    add_title_page(document)
    parse_markdown(document, lines, base_dir=os.path.dirname(os.path.abspath(src)))
    document.save(out)
    print(f"Zapisano: {out}")


if __name__ == "__main__":
    main()
